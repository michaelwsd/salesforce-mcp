"""Build an AA Investment Screener (.docx) from a JSON spec.

Usage:
    python build_screener.py spec.json                  # -> ./Screener_Project X.docx
    python build_screener.py spec.json -o out_dir/      # -> out_dir/Screener_Project X.docx
    python build_screener.py spec.json -o file.docx
    python build_screener.py spec.json --check          # validate only, no build

The spec format is documented in SKILL.md and examples/example_spec.json.
House rules (number formats, fixed thesis/Porter rows, revenue-stream style,
percentages summing to 100%) are enforced here; errors stop the build and
warnings are printed. The file is built and repaired in a temporary directory
and only copied to the destination when complete, so writing to a synced
folder (OneDrive) cannot leave a half-written, corrupt file.
"""

import argparse
import json
import os
import re
import shutil
import sys
import tempfile

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import itertools

from fix_xml import repair_docx, repair_tree

# ── AA house format (do not change; see SKILL.md "AA format reference") ──────
PAGE_W, PAGE_H = 11906, 16838                  # A4, twips
MARGIN_TOP, MARGIN_BOTTOM, MARGIN_LR = 567, 568, 1134
CONTENT_W = 9628                               # DXA
TEAL = "156082"                                # Accent1: category headers, rule
NAVY = "0E2841"                                # dk2: title, chart annotations
SECTION_FILL = "F2F2F2"
RECOMMENDATION = "9C0006"
GREY_TEXT = "595959"
BOX_BORDER = "BFBFBF"
FONT = "Calibri"
BODY_PT = 9
THESIS_COLS = (2830, 738, 6060)
PORTER_COLS = (2543, 603, 6482)
SOURCE_COLS = (2819, 6809)
CHART_COLS = (4900, 2150, 2458)
CHART_IN, DONUT_IN = 3.40, 1.49
CHART_H_IN = 2.30

# Chart colours
C_REV = "#156082"
C_EARN = "#C00000"
C_GM = "#70AD47"
C_NAVY = "#0E2841"
DONUT_COLOURS = ["#156082", "#2E75B6", "#9DC3E6", "#C9E2F5", "#7F7F7F", "#BFBFBF"]

# Fixed thesis structure (from current AA screeners). Keys are the spec keys.
THESIS = [
    (("Growing, high-margin and recurring revenues from diversified mission-critical "
      "products and services"), [
         ("growing_high_margin", "Growing, high margin"),
         ("recurring_revenue", "Recurring revenue"),
         ("differentiated", "Differentiated service/product with high barriers to exit"),
     ]),
    ("Large and growing fragmented market with high barriers to entry", [
        ("large_growing_market", "Large and growing market with barriers to entry"),
        ("fragmented_end_market", "Fragmented end market"),
    ]),
]
PORTER = [
    ("supplier_power", "Supplier power"),
    ("buyer_power", "Buyer power"),
    ("competitive_rivalry", "Competitive rivalry"),
    ("threat_of_substitutes", "Threat of substitutes"),
    ("threat_of_new_entrants", "Threat of new entrants"),
]
THESIS_RATINGS = {"Y", "N", "TBC", "Y/TBC", "N/TBC"}
PORTER_RATINGS = {"L", "M", "H", "L/M", "M/H"}
VERDICTS = ("Pass", "Proceed", "Further diligence")
EARNINGS_METRICS = {"EBITDA", "Adj. EBITDA", "EBIT", "Adj. EBIT"}

FULL_YEAR_ACTUAL = re.compile(r"^FY\d{2}A$")

# Places Calibri ships outside the system font folders (Word for Mac bundles it).
FONT_DIRS = [
    "/Applications/Microsoft Word.app/Contents/Resources/DFonts",
    "/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts",
    os.path.expandvars(r"%WINDIR%\Fonts"),
]


# ═════════════════════════════════════════════════════════════════════════════
# Number formatting (one decimal place everywhere, brackets for negatives)
# ═════════════════════════════════════════════════════════════════════════════

def fmt_m(v: float) -> str:
    return f"({abs(v):.1f})" if v < 0 else f"{v:.1f}"


def fmt_pct(v: float, signed: bool = False) -> str:
    if v < 0:
        return f"({abs(v):.1f}%)"
    return f"+{v:.1f}%" if signed else f"{v:.1f}%"


# ═════════════════════════════════════════════════════════════════════════════
# Validation
# ═════════════════════════════════════════════════════════════════════════════

class Problems:
    def __init__(self):
        self.errors: list[str] = []
        self.warnings: list[str] = []

    def error(self, msg):
        self.errors.append(msg)

    def warn(self, msg):
        self.warnings.append(msg)


STAFF_COUNT = re.compile(r"\b\d[\d,.]*\s*(staff|FTEs?|employees|people|heads?)\b", re.IGNORECASE)
TWO_DP_PCT = re.compile(r"\d+\.\d{2,}\s?%")
WHOLE_PCT = re.compile(r"(?<![\d.~])\d+%")
THOUSANDS = re.compile(r"\$\s?\d{1,3}(,\d{3})+(?![\d.]*\s?[mbk])", re.IGNORECASE)


ABBREVIATIONS = ("adj.", "approx.", "incl.", "excl.", "c.", "e.g.", "i.e.", "vs.", "no.", "est.", "avg.")


def _sentences(text: str) -> int:
    """Count sentences, not splitting after abbreviations like 'adj.'."""
    pieces = [s for s in re.split(r"(?<=[.!?])\s+(?=[A-Z(])", text.strip()) if s]
    return sum(1 for i in range(len(pieces))
               if i == 0 or not pieces[i - 1].lower().endswith(ABBREVIATIONS))


def _check_text(path: str, text: str, p: Problems):
    if TWO_DP_PCT.search(text):
        p.error(f"{path}: percentages must be one decimal place: {TWO_DP_PCT.search(text).group()}")
    m = WHOLE_PCT.search(text)
    if m and m.group() != "100%":
        p.warn(f"{path}: whole-number percentage '{m.group()}' - house style is one decimal place "
               "(use ~ for an approximate IM claim, e.g. ~99%)")
    if THOUSANDS.search(text):
        p.error(f"{path}: dollar values must be in $m, not thousands: {THOUSANDS.search(text).group()}")


def _walk_text(node, path, p):
    if isinstance(node, str):
        _check_text(path, node, p)
    elif isinstance(node, dict):
        for k, v in node.items():
            _walk_text(v, f"{path}.{k}", p)
    elif isinstance(node, list):
        for i, v in enumerate(node):
            _walk_text(v, f"{path}[{i}]", p)


def _require(spec, key, p, kind=str):
    v = spec.get(key)
    if v is None or (isinstance(v, str) and not v.strip()):
        p.error(f"missing required field '{key}'")
        return None
    if not isinstance(v, kind):
        p.error(f"'{key}' must be {kind.__name__}")
        return None
    return v


def validate(spec: dict) -> Problems:
    p = Problems()

    for key in ("project_name", "company_name", "author_initials", "date",
                "lead_para", "revenue_streams_intro"):
        _require(spec, key, p)
    name = spec.get("project_name") or ""
    if name and not name.upper().startswith("PROJECT "):
        p.warn("project_name should be the AA codename, e.g. 'PROJECT BUNDABERG'")

    # Revenue streams
    streams = _require(spec, "revenue_streams", p, list) or []
    if streams:
        total = 0.0
        for i, s in enumerate(streams):
            for f in ("name", "pct", "value_m", "description"):
                if s.get(f) in (None, ""):
                    p.error(f"revenue_streams[{i}] missing '{f}'")
            if isinstance(s.get("pct"), str) or isinstance(s.get("value_m"), str):
                p.error(f"revenue_streams[{i}]: pct and value_m must be numbers, not strings")
                continue
            total += s.get("pct") or 0
            for f in ("pct", "value_m"):
                if isinstance(s.get(f), (int, float)) and round(s[f], 1) != s[f]:
                    p.error(f"revenue_streams[{i}].{f} must be to one decimal place, got {s[f]}")
            desc = s.get("description") or ""
            if STAFF_COUNT.search(desc):
                p.error(f"revenue_streams[{i}]: no staff counts in revenue-stream bullets "
                        f"('{STAFF_COUNT.search(desc).group()}') - put FTE in lead_para")
        if abs(total - 100) > 0.15:
            p.error(f"revenue_streams percentages sum to {total:.1f}%, must be 100.0%")

    # Business items
    for i, item in enumerate(spec.get("business_items") or []):
        if not item.get("label") or not item.get("text"):
            p.error(f"business_items[{i}] needs 'label' and 'text'")
    labels = [i.get("label", "").lower() for i in spec.get("business_items") or []]
    for needed in ("revenue model", "go-to-market"):
        if needed not in labels:
            p.warn(f"business_items has no '{needed.capitalize()}' item")

    # Financials
    fin = _require(spec, "financials", p, dict) or {}
    if fin:
        periods = fin.get("periods") or []
        n = len(periods)
        if n < 2:
            p.error("financials.periods needs at least two periods")
        for key in ("revenue", "earnings"):
            series = fin.get(key)
            if not isinstance(series, list) or len(series) != n:
                p.error(f"financials.{key} must be a list the same length as periods ({n})")
            elif any(not isinstance(v, (int, float)) for v in series):
                p.error(f"financials.{key} must be numbers in A$m (no nulls or strings)")
        gm = fin.get("gross_margin_pct")
        if gm is not None and (not isinstance(gm, list) or len(gm) != n
                               or any(not isinstance(v, (int, float)) for v in gm)):
            p.error("financials.gross_margin_pct must be null or a list of numbers, same length as periods")
        if fin.get("earnings_metric", "EBITDA") not in EARNINGS_METRICS:
            p.error(f"financials.earnings_metric must be one of {sorted(EARNINGS_METRICS)}")
        if sum(bool(FULL_YEAR_ACTUAL.match(x)) for x in periods) < 2:
            p.warn("fewer than two full-year actual periods (FYxxA) - no growth arrows or CAGR")
        for x in periods:
            if not re.search(r"[AF]$", x):
                p.error(f"period '{x}' must end in A (actual) or F (forecast)")
        callout = fin.get("callout")
        if callout and callout.get("period") not in periods:
            p.error("financials.callout.period must be one of financials.periods")
        if any(isinstance(v, (int, float)) and v > 500 for v in fin.get("revenue") or []):
            p.error("financials.revenue looks like thousands - values must be A$m")

    # Revenue mix (donut)
    mix = _require(spec, "revenue_mix", p, list) or []
    if mix:
        if len(mix) > len(DONUT_COLOURS):
            p.error(f"revenue_mix supports at most {len(DONUT_COLOURS)} segments")
        for m in mix:
            if isinstance(m.get("pct"), (int, float)) and round(m["pct"], 1) != m["pct"]:
                p.error(f"revenue_mix '{m.get('label')}' pct must be to one decimal place, got {m['pct']}")
        if any(not isinstance(m.get("pct"), (int, float)) for m in mix):
            p.error("revenue_mix pct values must be numbers (47.1), not strings ('47%')")
        else:
            total = sum(m["pct"] for m in mix)
            if abs(total - 100) > 0.15:
                p.error(f"revenue_mix percentages sum to {total:.1f}%, must be 100.0%")
        for m in mix:
            if len(m.get("label", "")) > 18:
                p.warn(f"revenue_mix label '{m.get('label')}' is long and may clip - one or two words")

    notes = spec.get("chart_notes") or []
    if not 4 <= len(notes) <= 6:
        p.warn(f"chart_notes has {len(notes)} bullets; house standard is 4 to 6")

    # Transaction dynamics
    td = _require(spec, "transaction", p, dict) or {}
    if td:
        for key in ("lead", "source", "other_niches"):
            if not td.get(key):
                p.error(f"transaction.{key} is required")
        rec = td.get("recommendation") or {}
        if rec.get("verdict") not in VERDICTS:
            p.error(f"transaction.recommendation.verdict must be one of {VERDICTS}")
        if not rec.get("rationale"):
            p.error("transaction.recommendation.rationale is required")
        if not rec.get("gates"):
            p.warn("transaction.recommendation.gates is empty - list the numbered next-step gates")
        if re.search(r"@\w", td.get("source", "")):
            p.warn("transaction.source looks like it contains an email address - firm and contact name only")

    # Thesis
    thesis = _require(spec, "thesis", p, dict) or {}
    expected = {k for _, rows in THESIS for k, _ in rows}
    for key in sorted(set(thesis) - expected):
        p.error(f"thesis has unknown criterion '{key}' - criteria are fixed: {sorted(expected)}")
    for key in sorted(expected):
        row = thesis.get(key)
        if not row:
            p.error(f"thesis.{key} is missing")
            continue
        if row.get("rating") not in THESIS_RATINGS:
            p.error(f"thesis.{key}.rating must be one of {sorted(THESIS_RATINGS)}")
        if not row.get("evidence"):
            p.error(f"thesis.{key}.evidence is required (use TBC with what needs validating)")
        elif _sentences(row["evidence"]) > 2:
            p.warn(f"thesis.{key}.evidence is more than two sentences")

    # Porter
    porter = _require(spec, "porter", p, dict) or {}
    expected = {k for k, _ in PORTER}
    for key in sorted(set(porter) - expected):
        p.error(f"porter has unknown force '{key}' - forces are fixed: {sorted(expected)}")
    for key in sorted(expected):
        row = porter.get(key)
        if not row:
            p.error(f"porter.{key} is missing")
            continue
        if row.get("rating") not in PORTER_RATINGS:
            p.error(f"porter.{key}.rating must be one of {sorted(PORTER_RATINGS)}")
        if not row.get("commentary"):
            p.error(f"porter.{key}.commentary is required")
        elif _sentences(row["commentary"]) > 2:
            p.warn(f"porter.{key}.commentary is more than two sentences")

    _walk_text({k: v for k, v in spec.items() if k != "financials"}, "spec", p)
    return p


# ═════════════════════════════════════════════════════════════════════════════
# Charts
# ═════════════════════════════════════════════════════════════════════════════

_chart_font = None


def chart_font() -> str:
    """Return the chart font family: Calibri if it can be found, else a metric-
    compatible or default fallback. Registers Office-bundled Calibri if needed."""
    global _chart_font
    if _chart_font:
        return _chart_font
    from matplotlib import font_manager

    def available():
        return {f.name for f in font_manager.fontManager.ttflist}

    names = available()
    if "Calibri" not in names:
        for d in FONT_DIRS:
            for fn in ("Calibri.ttf", "Calibrib.ttf", "Calibrii.ttf", "Calibriz.ttf"):
                path = os.path.join(d, fn)
                if os.path.exists(path):
                    font_manager.fontManager.addfont(path)
        names = available()
    for family in ("Calibri", "Carlito"):
        if family in names:
            _chart_font = family
            break
    else:
        _chart_font = "DejaVu Sans"
        print("WARNING: Calibri not found; charts use DejaVu Sans. Install Calibri or Carlito "
              "for house-style charts.")
    return _chart_font

def growth_points(periods):
    """Indices of full-year actual periods, used for YoY arrows and CAGR."""
    return [i for i, x in enumerate(periods) if FULL_YEAR_ACTUAL.match(x)]


def _rect_overlap(a, b):
    return a[0] < b[2] and b[0] < a[2] and a[1] < b[3] and b[1] < a[3]


def render_combo(fin: dict, path: str) -> None:
    periods = fin["periods"]
    rev = [float(v) for v in fin["revenue"]]
    earn = [float(v) for v in fin["earnings"]]
    gm = fin.get("gross_margin_pct")
    metric = fin.get("earnings_metric", "EBITDA")
    n = len(periods)
    x = np.arange(n)
    w = 0.36
    is_fcst = [p.endswith("F") for p in periods]

    plt.rcParams.update({"font.family": chart_font(), "font.size": 6})
    fig, ax = plt.subplots(figsize=(CHART_IN, CHART_H_IN), dpi=300)

    R = max(max(rev), max(earn), 0.1)
    neg_floor = min(0.0, min(earn), min(rev))
    alpha = [0.45 if f else 1.0 for f in is_fcst]

    for i in range(n):
        ax.bar(x[i] - w / 2, rev[i], w, color=C_REV, alpha=alpha[i], zorder=3, linewidth=0)
        ax.bar(x[i] + w / 2, earn[i], w, color=C_EARN, alpha=alpha[i], zorder=3, linewidth=0)

    lab_h = 0.075 * R          # approx. height of a 6pt label in data units
    obstacles = []             # (x0, y0, x1, y1) rectangles in data coords

    for i in range(n):
        y = max(rev[i], 0) + 0.015 * R
        ax.text(x[i] - w / 2, y, fmt_m(rev[i]), ha="center", va="bottom", fontsize=5.6,
                color=C_NAVY, fontweight="bold", zorder=6)
        obstacles.append((x[i] - w / 2 - 0.22, y, x[i] - w / 2 + 0.22, y + lab_h))
        if earn[i] >= 0:
            y = earn[i] + 0.015 * R
            ax.text(x[i] + w / 2, y, fmt_m(earn[i]), ha="center", va="bottom", fontsize=5.6,
                    color=C_EARN, fontweight="bold", zorder=6)
            obstacles.append((x[i] + w / 2 - 0.22, y, x[i] + w / 2 + 0.22, y + lab_h))
        else:
            y = earn[i] - 0.015 * R
            ax.text(x[i] + w / 2, y, fmt_m(earn[i]), ha="center", va="top", fontsize=5.6,
                    color=C_EARN, fontweight="bold", zorder=6)

    # ── YoY growth arrows between consecutive full-year actuals ──────────────
    gp = growth_points(periods)
    pairs = [(a, b) for a, b in itertools.pairwise(gp) if b == a + 1 and rev[a] > 0]
    yoy_tops = {}
    for a, b in pairs:
        pct = (rev[b] - rev[a]) / rev[a] * 100
        xa, xb = x[a] - w / 2, x[b] - w / 2
        ya = max(rev[a], 0) + 0.17 * R
        yb = max(rev[b], 0) + 0.17 * R
        ax.annotate("", xy=(xb, yb), xytext=(xa, ya),
                    arrowprops={"arrowstyle": "-|>", "color": C_NAVY, "lw": 0.6, "mutation_scale": 5},
                    zorder=6, annotation_clip=False)
        xm, ym = (xa + xb) / 2, (ya + yb) / 2 + 0.05 * R
        ax.text(xm, ym, fmt_pct(pct, signed=True), ha="center", va="bottom", fontsize=5.4,
                fontweight="bold", color=C_NAVY, zorder=7,
                bbox={"boxstyle": "round,pad=0.22", "fc": "white", "ec": C_NAVY, "lw": 0.45})
        top = ym + 0.13 * R
        yoy_tops[(a, b)] = (xm, top)
        obstacles.append((xm - 0.3, ym - 0.02 * R, xm + 0.3, top))
        for t in np.linspace(0, 1, 7):
            px, py = xa + t * (xb - xa), ya + t * (yb - ya)
            obstacles.append((px - 0.04, py - 0.02 * R, px + 0.04, py + 0.02 * R))

    # ── CAGR across full-year actuals, parallel to the revenue slope ─────────
    highest = max([max(rev) + 0.12 * R] + [t for _, t in yoy_tops.values()])
    cagr_top = highest
    if len(gp) >= 2 and rev[gp[0]] > 0 and rev[gp[-1]] > 0:
        a, b = gp[0], gp[-1]
        fy_a, fy_b = int(periods[a][2:4]), int(periods[b][2:4])
        years = fy_b - fy_a
        if years >= 1:
            cagr = ((rev[b] / rev[a]) ** (1 / years) - 1) * 100
            xa, xb = x[a] - w / 2, x[b] - w / 2

            def base(px):
                return rev[a] + (px - xa) / (xb - xa) * (rev[b] - rev[a]) if xb != xa else rev[a]

            clear = [t - base(xm) for xm, t in yoy_tops.values()]
            clear += [max(rev[i], 0) + 0.12 * R - base(x[i] - w / 2) for i in range(a, b + 1)]
            off = max(clear) + 0.10 * R
            ya, yb = rev[a] + off, rev[b] + off
            ax.annotate("", xy=(xb, yb), xytext=(xa, ya),
                        arrowprops={"arrowstyle": "-|>", "color": C_NAVY, "lw": 0.75, "mutation_scale": 6},
                        zorder=6, annotation_clip=False)
            xm, ym = (xa + xb) / 2, (ya + yb) / 2 + 0.05 * R
            ax.text(xm, ym, f"CAGR {fmt_pct(cagr, signed=True)}", ha="center", va="bottom",
                    fontsize=5.6, fontweight="bold", color=C_NAVY, zorder=7,
                    bbox={"boxstyle": "round,pad=0.25", "fc": "white", "ec": C_NAVY, "lw": 0.5})
            cagr_top = max(ya, yb, ym + 0.14 * R)
            for t in np.linspace(0, 1, 9):
                px, py = xa + t * (xb - xa), ya + t * (yb - ya)
                obstacles.append((px - 0.05, py - 0.02 * R, px + 0.05, py + 0.02 * R))
            obstacles.append((xm - 0.45, ym - 0.02 * R, xm + 0.45, ym + 0.14 * R))

    # ── Callout for an unusual period (optional) ─────────────────────────────
    top = cagr_top
    callout = fin.get("callout")
    if callout:
        i = periods.index(callout["period"])
        ty = top + 0.22 * R
        # Vertical arrow down the revenue bar's centre line (between the YoY labels),
        # stopping above the YoY arrow junction; box clamped inside the plot.
        half = max(len(line) for line in callout["text"].split("\n")) * 0.045 + 0.1
        tx = min(max(x[i] - w / 2, -0.5 + half), n - 0.5 - half)
        ax.annotate(callout["text"], xy=(x[i] - w / 2, max(rev[i], 0) + 0.23 * R),
                    xytext=(tx, ty),
                    fontsize=5.2, color=C_NAVY, ha="center", va="bottom", zorder=8,
                    bbox={"boxstyle": "round,pad=0.3", "fc": "#EBF0FA", "ec": C_NAVY, "lw": 0.5},
                    arrowprops={"arrowstyle": "-|>", "color": C_NAVY, "lw": 0.5, "mutation_scale": 5})
        top = ty + 0.22 * R

    # ── Gross margin % line: pick the band with the fewest collisions ────────
    if gm:
        gm = [float(v) for v in gm]
        lo, hi = min(gm), max(gm)
        bands = [(0.42, 0.58), (0.58, 0.74), (0.26, 0.42), (top / R + 0.04, top / R + 0.20)]
        best, best_score = None, None
        for b0, b1 in bands:
            ys = [((b0 + b1) / 2 if hi == lo else b0 + (v - lo) / (hi - lo) * (b1 - b0)) * R for v in gm]
            score = 0
            for i, y in enumerate(ys):
                box = (x[i] - 0.28, y + 0.02 * R, x[i] + 0.28, y + 0.02 * R + lab_h * 1.25)
                score += sum(_rect_overlap(box, o) for o in obstacles) * 3
                for t in np.linspace(0, 1, 6)[:-1] if i < n - 1 else []:
                    px, py = x[i] + t, y + t * (ys[i + 1] - y)
                    score += sum(_rect_overlap((px - 0.02, py - 0.01 * R, px + 0.02, py + 0.01 * R), o)
                                 for o in obstacles)
            if best_score is None or score < best_score:
                best, best_score = ys, score
        ax.plot(x, best, color=C_GM, marker="o", lw=1.1, ms=2.2, zorder=5)
        for i, y in enumerate(best):
            neighbours = [best[j] for j in (i - 1, i + 1) if 0 <= j < n]
            below = bool(neighbours) and all(nb > y for nb in neighbours)
            ax.text(x[i], y - 0.04 * R if below else y + 0.04 * R, fmt_pct(gm[i]), ha="center",
                    va="top" if below else "bottom", fontsize=5.2, color=C_GM, fontweight="bold",
                    zorder=7, bbox={"boxstyle": "round,pad=0.12", "fc": "white", "ec": "none", "alpha": 0.9})
        top = max(top, max(best) + 0.035 * R + lab_h * 1.3)

    # ── Axes, margin row and legend ──────────────────────────────────────────
    ax.set_ylim(neg_floor * 1.3 - 0.02 * R if neg_floor < 0 else 0, top + 0.03 * R)
    ax.set_xlim(-0.5, n - 0.5)
    ax.set_xticks(x)
    ax.set_xticklabels(periods, fontsize=6, fontweight="bold", color=C_NAVY)
    ax.tick_params(axis="x", length=0, pad=2)
    ax.yaxis.set_visible(False)
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    ax.spines["bottom"].set_color("#A6A6A6")
    ax.spines["bottom"].set_linewidth(0.5)
    if neg_floor < 0:
        ax.spines["bottom"].set_visible(False)
        ax.axhline(0, color="#A6A6A6", lw=0.5, zorder=2)
        ax.tick_params(axis="x", pad=2 + 9)

    trans = ax.get_xaxis_transform()
    margin_y = -0.13 if neg_floor >= 0 else -0.16
    for i in range(n):
        if rev[i]:
            ax.text(x[i], margin_y, fmt_pct(earn[i] / rev[i] * 100), transform=trans, ha="center",
                    va="top", fontsize=5.2, color=C_EARN, fontstyle="italic")
    ax.text(-0.01, margin_y, f"{metric} %", transform=ax.transAxes, ha="right", va="top",
            fontsize=5.2, color=C_EARN, fontstyle="italic")

    from matplotlib.lines import Line2D
    from matplotlib.patches import Patch
    handles = [Patch(color=C_REV, label="Revenue"), Patch(color=C_EARN, label=metric)]
    if gm:
        handles.append(Line2D([0], [0], color=C_GM, marker="o", lw=1.1, ms=2.2, label="GP margin %"))
    if any(is_fcst):
        handles.append(Patch(color=C_REV, alpha=0.45, label="Forecast"))
    ax.legend(handles=handles, loc="upper center", bbox_to_anchor=(0.5, margin_y - 0.09),
              ncol=len(handles), fontsize=5.6, frameon=False, handlelength=1.2,
              handletextpad=0.4, columnspacing=1.0)

    fig.subplots_adjust(left=0.155, right=0.995, top=0.99, bottom=0.25)
    fig.savefig(path, dpi=300, facecolor="white")
    plt.close(fig)


def render_donut(mix: list[dict], centre: str, path: str) -> None:
    plt.rcParams.update({"font.family": chart_font()})
    fig = plt.figure(figsize=(DONUT_IN, CHART_H_IN), dpi=300)
    ax = fig.add_axes([0.04, 0.40, 0.92, 0.58])
    vals = [m["pct"] for m in mix]
    colours = DONUT_COLOURS[: len(mix)]
    ax.pie(vals, colors=colours, startangle=90, counterclock=False,
           wedgeprops={"width": 0.36, "edgecolor": "white", "linewidth": 0.8})
    ax.text(0, 0, centre, ha="center", va="center", fontsize=6, fontweight="bold", color=C_NAVY)
    ax.set_aspect("equal")
    ax.axis("off")

    from matplotlib.patches import Patch
    handles = [Patch(color=c, label=f"{m['label']} {fmt_pct(m['pct'])}") for m, c in zip(mix, colours)]
    fig.legend(handles=handles, loc="upper left", bbox_to_anchor=(0.03, 0.39), ncol=1,
               fontsize=5.4, frameon=False, handlelength=0.9, handleheight=0.9,
               handletextpad=0.4, labelspacing=0.35, borderaxespad=0)
    fig.savefig(path, dpi=300, facecolor="white")
    plt.close(fig)


# ═════════════════════════════════════════════════════════════════════════════
# Document helpers
# ═════════════════════════════════════════════════════════════════════════════

def _el(tag: str, **attrs) -> OxmlElement:
    e = OxmlElement(f"w:{tag}")
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    return e


def _replace(parent, child):
    for old in parent.findall(child.tag):
        parent.remove(old)
    parent.append(child)


def add_run(p, text, bold=False, italic=False, size=BODY_PT, color=None):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def tight(p, before=0, after=0):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0
    return p


def spacer(doc, pts):
    """Empty paragraph of a fixed height. Also stops Word merging adjacent tables."""
    p = tight(doc.add_paragraph())
    p.paragraph_format.line_spacing = Pt(pts)
    p.paragraph_format.keep_with_next = True
    add_run(p, "", size=1)
    return p


def section_heading(doc, text):
    """Grey bar with a teal rule, drawn as a one-cell table so its edges line up
    exactly with the tables (paragraph shading extends past the margins by an
    amount that depends on Word's border rendering)."""
    spacer(doc, 6)
    t = doc.add_table(rows=1, cols=1)
    configure_table(t, [CONTENT_W], borders=False)
    b = t._tbl.tblPr.find(qn("w:tblBorders"))
    b.remove(b.find(qn("w:top")))
    b.insert(0, _el("top", val="single", sz=8, space=0, color=TEAL))
    cell = t.rows[0].cells[0]
    shade(cell, SECTION_FILL)
    p = cell_text(cell, text, bold=True, size=10, color=NAVY)
    cell_margins(cell, top=15, bottom=15, left=60, right=60)
    p.paragraph_format.keep_with_next = True
    spacer(doc, 3)
    return p


def bullet(container, level=1):
    style = "List Bullet" if level == 1 else "List Bullet 2"
    return tight(container.add_paragraph(style=style))


def shade(cell, fill):
    _replace(cell._tc.get_or_add_tcPr(), _el("shd", val="clear", color="auto", fill=fill))


def cell_margins(cell, top=30, bottom=30, left=60, right=60):
    mar = _el("tcMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        mar.append(_el(side, w=val, type="dxa"))
    _replace(cell._tc.get_or_add_tcPr(), mar)


def cell_text(cell, text, bold=False, size=BODY_PT, color=None, align=None):
    cell.text = ""
    p = tight(cell.paragraphs[0])
    if align:
        p.alignment = align
    add_run(p, text, bold=bold, size=size, color=color)
    cell_margins(cell)
    return p


def configure_table(table, widths, borders=True, border_color="000000"):
    """Fixed layout, exact column widths, zero indent so the table is flush with
    the section headings."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    _replace(tblPr, _el("tblW", w=sum(widths), type="dxa"))
    _replace(tblPr, _el("tblInd", w=0, type="dxa"))
    _replace(tblPr, _el("tblLayout", type="fixed"))
    b = _el("tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if borders:
            b.append(_el(side, val="single", sz=4, space=0, color=border_color))
        else:
            b.append(_el(side, val="nil"))
    _replace(tblPr, b)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    grid = tbl.tblGrid
    for gc in list(grid):
        grid.remove(gc)
    for wdt in widths:
        grid.append(_el("gridCol", w=wdt))
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        _replace(trPr, _el("cantSplit"))
        for i, wdt in enumerate(widths):
            if i < len(row.cells):
                row.cells[i].width = Twips(wdt)


def repeat_header(row):
    _replace(row._tr.get_or_add_trPr(), _el("tblHeader"))


def merge_row(row, text, fill, color="FFFFFF"):
    merged = row.cells[0].merge(row.cells[-1])
    shade(merged, fill)
    cell_text(merged, text, bold=True, color=color)
    return merged


def add_field(p, instr, size, color):
    """Complex field (PAGE, NUMPAGES) whose result keeps the run formatting."""
    for kind in ("begin", "instr", "separate", "result", "end"):
        r = add_run(p, "", size=size, color=color)
        if kind == "instr":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {instr} "
            r._r.append(it)
        elif kind == "result":
            r.text = "1"
        else:
            r._r.append(_el("fldChar", fldCharType=kind))


def footer(doc, section, source):
    ft = section.footer
    p = tight(ft.paragraphs[0])
    p.style = doc.styles["Normal"]
    pPr = p._p.get_or_add_pPr()
    tabs = _el("tabs")
    tabs.append(_el("tab", val="right", pos=CONTENT_W))
    _replace(pPr, tabs)
    if source:
        add_run(p, source, size=7, color=GREY_TEXT)
    add_run(p, "\tPage ", size=7, color=GREY_TEXT)
    add_field(p, "PAGE", 7, GREY_TEXT)
    add_run(p, " of ", size=7, color=GREY_TEXT)
    add_field(p, "NUMPAGES", 7, GREY_TEXT)


def set_sub_bullet_glyph(doc):
    """Second-level bullets render as 'o' (Courier New), matching AA screeners."""
    style = doc.styles["List Bullet 2"].element
    num_id = style.pPr.find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))
    numbering = doc.part.numbering_part.element
    abstract_id = next(n.find(qn("w:abstractNumId")).get(qn("w:val"))
                       for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == num_id)
    abstract = next(a for a in numbering.findall(qn("w:abstractNum"))
                    if a.get(qn("w:abstractNumId")) == abstract_id)
    lvl = next(lv for lv in abstract.findall(qn("w:lvl")) if lv.get(qn("w:ilvl")) == "0")
    lvl.find(qn("w:lvlText")).set(qn("w:val"), "o")
    rPr = lvl.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        lvl.append(rPr)
    _replace(rPr, _el("rFonts", ascii="Courier New", hAnsi="Courier New", hint="default"))


def set_modern_compat(doc):
    """Word 2013+ compatibility mode, so table borders align with the margins."""
    settings = doc.settings.element
    for cs in settings.iter(qn("w:compatSetting")):
        if cs.get(qn("w:name")) == "compatibilityMode":
            cs.set(qn("w:val"), "15")


# ═════════════════════════════════════════════════════════════════════════════
# Document build
# ═════════════════════════════════════════════════════════════════════════════

def project_title(spec) -> str:
    """'PROJECT BUNDABERG' -> 'Project Bundaberg' (used in the file name)."""
    return spec["project_name"].strip().title()


def build(spec: dict, out_path: str, workdir: str) -> None:
    doc = Document()
    set_modern_compat(doc)
    set_sub_bullet_glyph(doc)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Twips(PAGE_W), Twips(PAGE_H)
    sec.top_margin, sec.bottom_margin = Twips(MARGIN_TOP), Twips(MARGIN_BOTTOM)
    sec.left_margin = sec.right_margin = Twips(MARGIN_LR)
    sec.header_distance = Twips(340)
    sec.footer_distance = Twips(284)
    footer(doc, sec, spec.get("source_note", ""))

    # ── Title ────────────────────────────────────────────────────────────────
    tp = tight(doc.add_paragraph())
    tabs = _el("tabs")
    tabs.append(_el("tab", val="right", pos=CONTENT_W))
    _replace(tp._p.get_or_add_pPr(), tabs)
    add_run(tp, f"SCREENING MEMO – {spec['project_name'].upper()} ({spec['company_name']})",
            bold=True, size=11, color=NAVY)
    add_run(tp, f"\t{spec['author_initials']}", size=9, color=GREY_TEXT)
    dp = tight(doc.add_paragraph(), after=2)
    add_run(dp, spec["date"], italic=True, size=9, color=GREY_TEXT)

    # ── Business & industry overview ─────────────────────────────────────────
    section_heading(doc, "Business & industry overview")
    add_run(tight(doc.add_paragraph(), after=2), spec["lead_para"])

    p = bullet(doc)
    add_run(p, "Revenue streams: ", bold=True)
    add_run(p, spec["revenue_streams_intro"])
    for s in spec["revenue_streams"]:
        p = bullet(doc, level=2)
        add_run(p, f"{s['name']} ({fmt_pct(s['pct'])}, ${s['value_m']:.1f}m)", bold=True)
        add_run(p, f" – {s['description']}")
    for item in spec.get("business_items") or []:
        p = bullet(doc)
        label = item["label"].rstrip(": ")
        add_run(p, f"{label}: ", bold=True)
        add_run(p, item["text"])

    # ── Financial overview ───────────────────────────────────────────────────
    fin = spec["financials"]
    section_heading(doc, "Financial overview – P&L and revenue segment mix")
    if spec.get("fin_note"):
        add_run(tight(doc.add_paragraph(), after=2), spec["fin_note"], italic=True, size=8,
                color=GREY_TEXT)

    combo_png = os.path.join(workdir, "combo.png")
    donut_png = os.path.join(workdir, "donut.png")
    render_combo(fin, combo_png)
    mix_period = spec.get("revenue_mix_period") or next(
        (x for x in reversed(fin["periods"]) if x.endswith("A")), fin["periods"][-1])
    i = fin["periods"].index(mix_period) if mix_period in fin["periods"] else None
    centre = f"${fin['revenue'][i]:.1f}m\n{mix_period}" if i is not None else mix_period
    render_donut(spec["revenue_mix"], centre, donut_png)

    box = doc.add_table(rows=1, cols=1)
    configure_table(box, [CONTENT_W], border_color=BOX_BORDER)
    outer = box.rows[0].cells[0]
    cell_margins(outer, top=40, bottom=40, left=60, right=60)
    inner = outer.add_table(rows=1, cols=3)
    configure_table(inner, list(CHART_COLS), borders=False)
    outer._tc.remove(outer.paragraphs[0]._p)       # drop the empty lead paragraph
    tight(outer.paragraphs[-1])                    # trailing paragraph Word requires
    outer.paragraphs[-1].paragraph_format.line_spacing = Pt(1)

    c_chart, c_donut, c_notes = inner.rows[0].cells
    metric = fin.get("earnings_metric", "EBITDA")
    for cell, caption in ((c_chart, f"Revenue & {metric} (A$m)"), (c_donut, f"Revenue mix {mix_period}"),
                          (c_notes, "Notes")):
        cell_margins(cell, top=0, bottom=0, left=40, right=40)
        cell_text(cell, caption, bold=True, size=8, color=NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    tight(c_chart.add_paragraph()).add_run().add_picture(combo_png, width=Inches(CHART_IN))
    tight(c_donut.add_paragraph()).add_run().add_picture(donut_png, width=Inches(DONUT_IN))
    for note in spec.get("chart_notes") or []:
        p = bullet(c_notes)
        p.paragraph_format.left_indent = Twips(170)
        p.paragraph_format.first_line_indent = Twips(-170)
        p.paragraph_format.space_after = Pt(1)
        add_run(p, note, size=7.5)

    # ── Transaction dynamics ─────────────────────────────────────────────────
    td = spec["transaction"]
    section_heading(doc, "Transaction dynamics, recommendation & next steps")
    add_run(tight(doc.add_paragraph(), after=3), td["lead"])

    rec = td["recommendation"]
    rec_text = f"{rec['verdict']} – {rec['rationale'].strip().rstrip('.')}"
    if rec.get("gates"):
        gates = "; ".join(f"({k}) {g.strip().rstrip(';.')}" for k, g in enumerate(rec["gates"], 1))
        rec_text += f". Gates: {gates}"
    rows = [("Source", td["source"], False), ("Next steps", rec_text, True),
            ("Other niches", td["other_niches"], False)]
    t = doc.add_table(rows=len(rows), cols=2)
    configure_table(t, list(SOURCE_COLS))
    for row, (label, text, is_rec) in zip(t.rows, rows):
        shade(row.cells[0], SECTION_FILL)
        cell_text(row.cells[0], label, bold=True)
        cell_text(row.cells[1], text, bold=is_rec, color=RECOMMENDATION if is_rec else None)

    # ── Investment thesis ────────────────────────────────────────────────────
    section_heading(doc, "Investment thesis criteria")
    n_rows = sum(1 + len(rows) for _, rows in THESIS)
    t = doc.add_table(rows=n_rows, cols=3)
    configure_table(t, list(THESIS_COLS))
    r = 0
    for header, criteria in THESIS:
        merged = merge_row(t.rows[r], header, TEAL)
        merged.paragraphs[0].paragraph_format.keep_with_next = True
        r += 1
        for key, label in criteria:
            row = t.rows[r]
            entry = spec["thesis"][key]
            cell_text(row.cells[0], label, bold=True)
            cell_text(row.cells[1], entry["rating"], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            cell_text(row.cells[2], entry["evidence"])
            r += 1

    # ── Porter's five forces ─────────────────────────────────────────────────
    section_heading(doc, "Porter's five forces")
    t = doc.add_table(rows=1 + len(PORTER), cols=3)
    configure_table(t, list(PORTER_COLS))
    hdr = t.rows[0]
    repeat_header(hdr)
    for cell, label in zip(hdr.cells, ("Force", "Rating", "Commentary")):
        shade(cell, SECTION_FILL)
        cell_text(cell, label, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER if label == "Rating" else None)
    for row, (key, label) in zip(t.rows[1:], PORTER):
        entry = spec["porter"][key]
        cell_text(row.cells[0], label, bold=True)
        cell_text(row.cells[1], entry["rating"], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(row.cells[2], entry["commentary"])

    # Schema-order every property element, then write and repair settings.
    repair_tree(doc.element.body)
    raw = os.path.join(workdir, "raw.docx")
    doc.save(raw)
    repair_docx(raw, out_path)


def resolve_output(spec: dict, out: str | None) -> str:
    name = f"Screener_{project_title(spec)}.docx"
    if not out:
        return os.path.abspath(name)
    if out.endswith(os.sep) or os.path.isdir(out):
        return os.path.abspath(os.path.join(out, name))
    return os.path.abspath(out)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec", help="Screener spec JSON file")
    ap.add_argument("-o", "--out", help="Output .docx path or directory")
    ap.add_argument("--check", action="store_true", help="Validate the spec only")
    ap.add_argument("--strict", action="store_true", help="Treat warnings as errors")
    args = ap.parse_args()

    with open(args.spec, encoding="utf-8") as f:
        spec = json.load(f)

    problems = validate(spec)
    for w in problems.warnings:
        print(f"WARNING: {w}")
    for e in problems.errors:
        print(f"ERROR: {e}")
    if problems.errors or (args.strict and problems.warnings):
        sys.exit(f"\n{len(problems.errors)} error(s), {len(problems.warnings)} warning(s) - not built.")
    if args.check:
        print(f"Spec OK ({len(problems.warnings)} warning(s)).")
        return

    dest = resolve_output(spec, args.out)
    with tempfile.TemporaryDirectory() as work:
        built = os.path.join(work, os.path.basename(dest))
        build(spec, built, work)
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        shutil.copyfile(built, dest)
    print(f"Built {dest}")


if __name__ == "__main__":
    main()
